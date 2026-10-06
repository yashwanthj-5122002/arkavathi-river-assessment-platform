from __future__ import annotations

import argparse
import json
import math
import re
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from openpyxl import load_workbook

K = "{http://www.opengis.net/kml/2.2}"
GX = "{http://www.google.com/kml/ext/2.2}"


def child_text(el: ET.Element | None, tag: str) -> str:
    if el is None:
        return ""
    c = el.find(K + tag)
    return c.text.strip() if c is not None and c.text else ""


def direct_children(el: ET.Element, tag: str) -> list[ET.Element]:
    return [c for c in list(el) if c.tag == K + tag]


def parse_extended(pm: ET.Element) -> dict[str, str]:
    data: dict[str, str] = {}
    ed = pm.find(K + "ExtendedData")
    if ed is None:
        return data
    for d in ed.findall(K + "Data"):
        name = d.get("name", "")
        v = d.find(K + "value")
        data[name] = v.text.strip() if v is not None and v.text else ""
    return data


def parse_coords(text: str) -> list[list[float | None]]:
    coords: list[list[float | None]] = []
    for token in (text or "").split():
        parts = token.strip().split(",")
        if len(parts) < 2:
            continue
        lon = float(parts[0])
        lat = float(parts[1])
        alt = float(parts[2]) if len(parts) > 2 and parts[2] else None
        coords.append([lon, lat, alt])
    return coords


def parse_flight_description(desc: str) -> dict:
    clean = re.sub(r"<br\s*/?>", " ", desc or "", flags=re.I)
    clean = re.sub(r"<[^>]+>", " ", clean)
    clean = clean.replace("&#183;", "·").replace("&middot;", "·")
    clean = re.sub(r"\s+", " ", clean).strip()
    m = re.search(
        r"(\d+)\s+points\s+·\s+([0-9.]+)\s+km\s+·\s+(.+?)\s+altitude\s+([0-9.]+)\s+-\s+([0-9.]+)\s+m",
        clean,
        re.I,
    )
    if not m:
        return {"raw": clean}
    return {
        "reported_points": int(m.group(1)),
        "reported_distance_km": float(m.group(2)),
        "reported_duration": m.group(3).strip(),
        "alt_min_m": float(m.group(4)),
        "alt_max_m": float(m.group(5)),
        "raw": clean,
    }


def to_float(value: str | None) -> float | None:
    try:
        return float(value) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def to_int(value: str | None) -> int | None:
    try:
        return int(float(value)) if value not in (None, "") else None
    except (TypeError, ValueError):
        return None


def haversine_m(a: list[float | None], b: list[float | None]) -> float:
    lon1, lat1 = float(a[0]), float(a[1])
    lon2, lat2 = float(b[0]), float(b[1])
    radius = 6_371_008.8
    p1, p2 = math.radians(lat1), math.radians(lat2)
    dp = math.radians(lat2 - lat1)
    dl = math.radians(lon2 - lon1)
    h = math.sin(dp / 2) ** 2 + math.cos(p1) * math.cos(p2) * math.sin(dl / 2) ** 2
    return 2 * radius * math.asin(min(1, math.sqrt(h)))


def parse_duration_label(seconds: int) -> str:
    h, rem = divmod(int(seconds), 3600)
    m, s = divmod(rem, 60)
    if h:
        return f"{h} h {m:02d} m {s:02d} s"
    if m:
        return f"{m} m {s:02d} s"
    return f"{s} s"


def duration_label_to_seconds(label: str | None) -> int:
    if not label or label.lower() in {"n/a", "na"}:
        return 0
    h = re.search(r"(\d+)\s*h", label, re.I)
    m = re.search(r"(\d+)\s*m", label, re.I)
    sec = re.search(r"(\d+)\s*s", label, re.I)
    return (int(h.group(1)) if h else 0) * 3600 + (int(m.group(1)) if m else 0) * 60 + (int(sec.group(1)) if sec else 0)

def parse_root_report(desc: str) -> dict:
    clean = desc or ""
    total = re.search(
        r"<b>(\d+)</b>\s*day\(s\).*?<b>(\d+)</b>\s*clip\(s\).*?<b>(\d+)</b>\s*track points from\s*(\d+)\s*source rows.*?<b>([0-9.]+) km</b> flown.*?<b>(.*?)</b> airborne combined",
        clean,
        flags=re.I | re.S,
    )
    report: dict = {}
    if total:
        report = {
            "survey_days": int(total.group(1)),
            "flight_clips": int(total.group(2)),
            "reported_track_points": int(total.group(3)),
            "source_rows": int(total.group(4)),
            "flight_distance_km": float(total.group(5)),
            "airborne_combined": re.sub(r"\s+", " ", total.group(6)).strip(),
        }
    day_rows = re.findall(
        r"<tr><td>Day\s+(\d+)</td><td[^>]*>(\d+)</td><td[^>]*>([0-9.]+) km</td><td[^>]*>(.*?)</td><td[^>]*>([0-9.]+) - ([0-9.]+) m</td></tr>",
        clean,
        flags=re.I | re.S,
    )
    report["days"] = {
        int(day): {
            "clips": int(clips),
            "distance_km": float(distance),
            "duration_label": re.sub(r"\s+", " ", duration).strip(),
            "alt_min_m": float(amin),
            "alt_max_m": float(amax),
        }
        for day, clips, distance, duration, amin, amax in day_rows
    }
    return report


def import_zip(zip_path: Path, output_dir: Path) -> dict:
    output_dir.mkdir(parents=True, exist_ok=True)
    with tempfile.TemporaryDirectory(prefix="arkavathi_import_") as tmp:
        tmp_path = Path(tmp)
        with zipfile.ZipFile(zip_path) as zf:
            zf.extractall(tmp_path)

        candidates = list(tmp_path.rglob("Merged.kmz"))
        if not candidates:
            raise FileNotFoundError("Merged.kmz not found inside the ZIP archive")
        merged_kmz = candidates[0]
        root_dir = merged_kmz.parent.parent

        workbook_info: dict[int, dict] = {}
        for day in range(3, 9):
            workbook = root_dir / f"Day_{day}.xlsx"
            if not workbook.exists():
                raise FileNotFoundError(f"Missing {workbook.name}")
            wb = load_workbook(workbook, read_only=True, data_only=True)
            ws = wb.active
            headers = [cell.value for cell in next(ws.iter_rows(min_row=1, max_row=1))]
            workbook_info[day] = {
                "source_rows": ws.max_row - 1,
                "columns": headers,
                "sheet": ws.title,
                "file": workbook.name,
            }
            wb.close()

        with zipfile.ZipFile(merged_kmz) as kmz:
            root = ET.fromstring(kmz.read("doc.kml"))
        document = root.find(K + "Document")
        if document is None:
            raise ValueError("Invalid KML: Document element is missing")
        root_report = parse_root_report(child_text(document, "description"))

        flights: list[dict] = []
        days: list[dict] = []
        all_coords: list[list[float | None]] = []
        issues: list[str] = []

        for day_folder in direct_children(document, "Folder"):
            m = re.fullmatch(r"Day\s+(\d+)", child_text(day_folder, "name"))
            if not m:
                continue
            day = int(m.group(1))
            if day < 3 or day > 8:
                continue

            animated_tracks: list[list[dict]] = []
            for pm in direct_children(day_folder, "Placemark"):
                if child_text(pm, "name") != "Animated Flight":
                    continue
                multi = pm.find(GX + "MultiTrack")
                if multi is None:
                    continue
                for track in multi.findall(GX + "Track"):
                    whens = [w.text.strip() for w in track.findall(K + "when") if w.text]
                    coords = []
                    for c in track.findall(GX + "coord"):
                        if not c.text:
                            continue
                        parts = c.text.split()
                        if len(parts) >= 3:
                            coords.append([float(parts[0]), float(parts[1]), float(parts[2])])
                    n = min(len(whens), len(coords))
                    animated_tracks.append(
                        [
                            {
                                "time": whens[i],
                                "lon": coords[i][0],
                                "lat": coords[i][1],
                                "abs_alt_m": coords[i][2],
                            }
                            for i in range(n)
                        ]
                    )

            clips_parent = next(
                (f for f in direct_children(day_folder, "Folder") if child_text(f, "name") == "Flight paths by clip"),
                None,
            )
            if clips_parent is None:
                issues.append(f"Day {day}: Flight paths by clip folder not found")
                continue

            day_flights: list[dict] = []
            for sequence, clip_folder in enumerate(direct_children(clips_parent, "Folder"), start=1):
                clip = child_text(clip_folder, "name")
                placemarks = {child_text(pm, "name"): pm for pm in direct_children(clip_folder, "Placemark")}
                path_pm = placemarks.get("Flight path")
                start_pm = placemarks.get("Start")
                end_pm = placemarks.get("End")
                if path_pm is None or start_pm is None or end_pm is None:
                    issues.append(f"Day {day} / {clip}: missing Flight path, Start or End placemark")
                    continue

                line = path_pm.find(K + "LineString")
                coords = parse_coords(child_text(line, "coordinates"))
                all_coords.extend(coords)
                desc = parse_flight_description(child_text(path_pm, "description"))
                start = parse_extended(start_pm)
                end = parse_extended(end_pm)
                path_len_m = to_float(end.get("Path length (m)"))
                recomputed_len = sum(haversine_m(a, b) for a, b in zip(coords, coords[1:]))
                distance_km = path_len_m / 1000 if path_len_m is not None else recomputed_len / 1000
                playback = animated_tracks[sequence - 1] if sequence - 1 < len(animated_tracks) else []

                flight = {
                    "id": f"D{day}-F{sequence:03d}",
                    "day": day,
                    "sequence": sequence,
                    "clip": clip,
                    "distance_km": round(distance_km, 4),
                    "reported_points": desc.get("reported_points"),
                    "rendered_vertices": len(coords),
                    "duration_label": end.get("Duration") or desc.get("reported_duration") or "n/a",
                    "alt_min_m": desc.get("alt_min_m"),
                    "alt_max_m": desc.get("alt_max_m"),
                    "start": {
                        "time": start.get("Time") or None,
                        "frame": to_int(start.get("Frame")),
                        "lat": to_float(start.get("Latitude")),
                        "lon": to_float(start.get("Longitude")),
                        "rel_alt_m": to_float(start.get("Rel_Alt (m)")),
                        "abs_alt_m": to_float(start.get("Abs_Alt (m)")),
                    },
                    "end": {
                        "time": end.get("Time") or None,
                        "frame": to_int(end.get("Frame")),
                        "lat": to_float(end.get("Latitude")),
                        "lon": to_float(end.get("Longitude")),
                        "rel_alt_m": to_float(end.get("Rel_Alt (m)")),
                        "abs_alt_m": to_float(end.get("Abs_Alt (m)")),
                    },
                    "coordinates": coords,
                    "playback": playback,
                }
                flights.append(flight)
                day_flights.append(flight)

            report_day = root_report.get("days", {}).get(day, {})
            source = workbook_info[day]
            date_values = sorted({f["start"]["time"][:10] for f in day_flights if f["start"]["time"]})
            duration_seconds = 0
            for flight in day_flights:
                if flight["start"]["time"] and flight["end"]["time"]:
                    try:
                        duration_seconds += round(
                            (datetime.fromisoformat(flight["end"]["time"]) - datetime.fromisoformat(flight["start"]["time"]))
                            .total_seconds()
                        )
                    except ValueError:
                        pass

            alts = [
                value
                for f in day_flights
                for value in (f["alt_min_m"], f["alt_max_m"])
                if value is not None
            ]
            reported_duration_label = report_day.get("duration_label", parse_duration_label(duration_seconds) if duration_seconds else "n/a")
            reported_duration_seconds = duration_label_to_seconds(reported_duration_label)
            day_item = {
                "day": day,
                "date": date_values[0] if len(date_values) == 1 else None,
                "clips": len(day_flights),
                "distance_km": report_day.get("distance_km", round(sum(f["distance_km"] for f in day_flights), 2)),
                "duration_seconds": reported_duration_seconds,
                "recomputed_start_end_seconds": duration_seconds,
                "duration_label": reported_duration_label,
                "duration_available": bool(duration_seconds),
                "alt_min_m": report_day.get("alt_min_m", round(min(alts), 1) if alts else None),
                "alt_max_m": report_day.get("alt_max_m", round(max(alts), 1) if alts else None),
                "source_rows": source["source_rows"],
                "columns": source["columns"],
                "workbook": source["file"],
                "reported_track_points": sum(f["reported_points"] or 0 for f in day_flights),
                "rendered_vertices": sum(f["rendered_vertices"] for f in day_flights),
            }
            days.append(day_item)

            if report_day and len(day_flights) != report_day.get("clips"):
                issues.append(
                    f"Day {day}: parsed {len(day_flights)} clips but Merged.kmz reports {report_day.get('clips')}"
                )
            if animated_tracks and len(animated_tracks) != len(day_flights):
                issues.append(
                    f"Day {day}: parsed {len(animated_tracks)} animated gx:Track objects but {len(day_flights)} flight clips"
                )

        if not all_coords:
            raise ValueError("No flight path coordinates were parsed from Merged.kmz")
        bbox = [
            min(float(c[0]) for c in all_coords),
            min(float(c[1]) for c in all_coords),
            max(float(c[0]) for c in all_coords),
            max(float(c[1]) for c in all_coords),
        ]

        reported = root_report or {}
        exact_path_distance = round(sum(f["distance_km"] for f in flights), 3)
        total_duration_seconds = duration_label_to_seconds(reported.get("airborne_combined")) or sum(d["duration_seconds"] for d in days)
        recomputed_duration_seconds = sum(d["recomputed_start_end_seconds"] for d in days)
        summary = {
            "project": "Arkavati River Population Assessment Project",
            "dataset": "Drone survey telemetry and KMZ projections (Day 3 to Day 8)",
            "survey_days": reported.get("survey_days", len(days)),
            "flight_clips": reported.get("flight_clips", len(flights)),
            "source_rows": reported.get("source_rows", sum(d["source_rows"] for d in days)),
            "reported_track_points": reported.get(
                "reported_track_points", sum(d["reported_track_points"] for d in days)
            ),
            "rendered_vertices": sum(d["rendered_vertices"] for d in days),
            "flight_distance_km": reported.get("flight_distance_km", round(exact_path_distance, 2)),
            "recomputed_path_distance_km": exact_path_distance,
            "recorded_duration_seconds": total_duration_seconds,
            "recomputed_start_end_seconds": recomputed_duration_seconds,
            "airborne_combined": reported.get("airborne_combined", parse_duration_label(total_duration_seconds)),
            "duration_note": "Day 3 has no Date/Time/FrameCnt columns, so duration is unavailable for Day 3 and excluded from the combined recorded duration.",
            "alt_min_m": min(d["alt_min_m"] for d in days if d["alt_min_m"] is not None),
            "alt_max_m": max(d["alt_max_m"] for d in days if d["alt_max_m"] is not None),
            "bbox": [round(v, 6) for v in bbox],
            "days": sorted(days, key=lambda item: item["day"]),
        }

        checks = {
            "source_rows_match_merged_kmz": summary["source_rows"] == sum(d["source_rows"] for d in days),
            "clip_count_matches_merged_kmz": summary["flight_clips"] == len(flights),
            "reported_track_points_match_flight_descriptions": summary["reported_track_points"]
            == sum(f["reported_points"] or 0 for f in flights),
            "combined_duration_difference_seconds": abs(summary["recorded_duration_seconds"] - recomputed_duration_seconds),
            "distance_difference_km": round(abs(summary["flight_distance_km"] - exact_path_distance), 3),
            "all_flights_have_start_end": all(f["start"]["lat"] is not None and f["end"]["lat"] is not None for f in flights),
        }
        if checks["combined_duration_difference_seconds"] > 5:
            issues.append(
                f"Recomputed start/end duration differs from Merged.kmz summary by {checks['combined_duration_difference_seconds']} seconds"
            )
        if checks["distance_difference_km"] > 0.02:
            issues.append(
                f"Recomputed path distance differs from Merged.kmz summary by {checks['distance_difference_km']} km"
            )
        for key in (
            "source_rows_match_merged_kmz",
            "clip_count_matches_merged_kmz",
            "reported_track_points_match_flight_descriptions",
            "all_flights_have_start_end",
        ):
            if not checks[key]:
                issues.append(f"Validation failed: {key}")

        map_features = [
            {
                "type": "Feature",
                "id": f["id"],
                "properties": {
                    "id": f["id"],
                    "day": f["day"],
                    "sequence": f["sequence"],
                    "clip": f["clip"],
                    "distance_km": f["distance_km"],
                    "duration": f["duration_label"],
                    "alt_min_m": f["alt_min_m"],
                    "alt_max_m": f["alt_max_m"],
                },
                "geometry": {"type": "LineString", "coordinates": f["coordinates"]},
            }
            for f in flights
        ]
        endpoints = []
        for f in flights:
            for kind in ("start", "end"):
                point = f[kind]
                endpoints.append(
                    {
                        "type": "Feature",
                        "properties": {
                            "flight_id": f["id"],
                            "day": f["day"],
                            "clip": f["clip"],
                            "kind": kind,
                            "label": "Start" if kind == "start" else "End",
                            "time": point["time"],
                            "frame": point["frame"],
                            "rel_alt_m": point["rel_alt_m"],
                            "abs_alt_m": point["abs_alt_m"],
                        },
                        "geometry": {
                            "type": "Point",
                            "coordinates": [point["lon"], point["lat"], point["abs_alt_m"]],
                        },
                    }
                )

        availability = {
            "available_now": [
                {"key": "drone_tracks", "label": "Drone flight paths", "status": "available", "source": "Merged.kmz"},
                {"key": "start_end", "label": "Start / end locations", "status": "available", "source": "Merged.kmz"},
                {"key": "video_clip_ids", "label": "Source video clip IDs", "status": "available", "source": "Day 3-8 XLSX / KMZ"},
                {"key": "gps", "label": "Latitude / longitude", "status": "available", "source": "Day 3-8 XLSX"},
                {"key": "absolute_altitude", "label": "Absolute altitude", "status": "available", "source": "Day 3-8 XLSX"},
                {"key": "relative_altitude", "label": "Relative altitude", "status": "partial", "source": "Day 4-8 XLSX only"},
                {"key": "time_frame", "label": "Date, time and frame count", "status": "partial", "source": "Day 4-8 XLSX only"},
            ],
            "needed_next": [
                {"module": "River context", "data": "Arkavathi river and tributary alignment (KML/KMZ/SHP/GeoJSON)"},
                {"module": "True survey progress", "data": "Planned survey reach/length or planned survey GIS line"},
                {"module": "12 institutions", "data": "Institution-to-watershed/reach allocation with boundaries if available"},
                {"module": "River survey", "data": "Location ID, chainage, width, depth, bed level, bank/flow condition"},
                {"module": "Structures", "data": "Bridge/culvert/check dam/weir inventory with GPS, dimensions and condition"},
                {"module": "Drains & outfalls", "data": "Drain/outfall IDs, GPS, type/source, connectivity and flow condition"},
                {"module": "Water sampling", "data": "Sample ID, GPS/location, institution, date, status and lab report link"},
                {"module": "Critical locations", "data": "Category, problem, priority, recommended action and status"},
            ],
        }

        verification = {
            "validation_passed": len(issues) == 0,
            "issues": issues,
            "checks": checks,
            "note": "Merged.kmz reports 17,923 track points. The rendered LineString geometry contains 17,794 coordinate vertices because the KML flight-path geometry omits/reduces some repeated sampled points. The dashboard preserves the reported 17,923 KPI and separately exposes rendered vertices for technical verification.",
        }

        (output_dir / "summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")
        (output_dir / "flights.json").write_text(json.dumps(flights, separators=(",", ":")), encoding="utf-8")
        (output_dir / "flight_tracks.geojson").write_text(
            json.dumps({"type": "FeatureCollection", "features": map_features}, separators=(",", ":")),
            encoding="utf-8",
        )
        (output_dir / "endpoints.geojson").write_text(
            json.dumps({"type": "FeatureCollection", "features": endpoints}, separators=(",", ":")),
            encoding="utf-8",
        )
        (output_dir / "data_availability.json").write_text(json.dumps(availability, indent=2), encoding="utf-8")
        (output_dir / "verification.json").write_text(json.dumps(verification, indent=2), encoding="utf-8")
        return verification


def main() -> None:
    parser = argparse.ArgumentParser(description="Import Arkavathi Day 3-Day 8 XLSX + KMZ ZIP into dashboard JSON")
    parser.add_argument("zip_path", type=Path, help="Path to Output(2).zip or equivalent archive")
    parser.add_argument(
        "--output-dir",
        type=Path,
        default=Path(__file__).resolve().parents[1] / "app" / "data",
        help="Dashboard data output directory",
    )
    args = parser.parse_args()
    resolved_zip = args.zip_path.resolve()
    resolved_output = args.output_dir.resolve()
    verification = import_zip(resolved_zip, resolved_output)
    # Build the Google-Earth-style clickable point layer from the same source XLSX files.
    from build_clickable_points import build as build_clickable_points
    point_report = build_clickable_points(resolved_zip, resolved_output)
    verification["clickable_points"] = point_report
    print(json.dumps(verification, indent=2))
    raise SystemExit(0 if verification["validation_passed"] else 2)


if __name__ == "__main__":
    main()
